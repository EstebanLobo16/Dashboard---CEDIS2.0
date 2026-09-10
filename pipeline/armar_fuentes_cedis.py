#!/usr/bin/env python3
"""
Arma el fuentes.json que consume pipeline/correr_motor.js, a partir de los
archivos que entrega el área.

Es el equivalente local de 20_Fuentes.gs: hace lo mismo que la ingesta de Apps
Script —concentrar los CSV en un padrón y una tabla de finalizaciones, leer las
cuatro pestañas de cada PDT, corregir el encabezado corrido del gerencial— pero
en la máquina, para poder correr el motor contra datos reales sin desplegar nada.

    python3 pipeline/armar_fuentes_cedis.py <carpeta-con-los-crudos> [fuentes.json]
    node pipeline/correr_motor.js fuentes.json

Espera encontrar en la carpeta:
    CEDIS P*.csv                    los cortes de finalizaciones
    PDT-operacion-*.xlsx            plan Colaborador
    PDT-gerencial-*.xlsx            plan Gerencial
    detalle_colaborador.xlsx        opcional
"""
import csv, glob, json, os, sys, unicodedata
from datetime import datetime, date, time, timedelta

csv.field_size_limit(10 ** 9)

NIVELES_MATRIZ = [
    'Jefes y Coordinadores', 'Gerente Operación', 'Gerente de Zona/Gte Sr',
    'Gerente Regional', 'Gerente Divisional o Director de Área',
    'Director Corporativo o Director General',
]


def clave(valor):
    """El textoClave_ de 10_Util.gs: sin acentos, mayúsculas, un solo espacio."""
    texto = '' if valor is None else str(valor)
    texto = texto.replace('\xa0', ' ')
    texto = unicodedata.normalize('NFD', texto)
    texto = ''.join(c for c in texto if unicodedata.category(c) != 'Mn')
    return ' '.join(texto.split()).upper()


def uno(carpeta, patron, obligatorio=True):
    encontrados = sorted(glob.glob(os.path.join(carpeta, patron)))
    if not encontrados:
        if obligatorio:
            sys.exit(f'No encontré ningún archivo que combine con "{patron}" en {carpeta}')
        return None
    return encontrados[-1]


def texto_celda(valor):
    """Lo que normalizarCelda_ deja: las fechas como AAAA-MM-DD, el resto tal cual."""
    if isinstance(valor, datetime):
        return valor.strftime('%Y-%m-%d')
    if isinstance(valor, date):
        return valor.strftime('%Y-%m-%d')
    if isinstance(valor, (time, timedelta)):
        return str(valor)
    if valor is None:
        return ''
    return valor.strip() if isinstance(valor, str) else valor


# --------------------------------------------------------------------------
#  Los CSV: un padrón y una tabla de finalizaciones
# --------------------------------------------------------------------------
COLUMNAS_PADRON = {
    'Número Persona': 'numeroPersona', 'Número Colaborador': 'numeroColaborador',
    'Nombre Colaborador': 'nombre', 'Región RRHH': 'region',
    'Centro Costos': 'centroCostos', 'Área': 'area', 'Departamento': 'departamento',
    'Puesto': 'puesto', 'Tipo Posición': 'categoria',
    'Fecha Contratación': 'fechaContratacion', 'Fecha Asignación Puesto': 'fechaPuesto',
}


def leer_csv(carpeta):
    padron, finalizaciones, vistos = {}, [], set()
    archivos = sorted(glob.glob(os.path.join(carpeta, 'CEDIS P*.csv')))
    if not archivos:
        sys.exit(f'No encontré ningún "CEDIS P*.csv" en {carpeta}')

    for ruta in archivos:
        filas = 0
        with open(ruta, newline='', encoding='utf-8-sig') as fh:
            for fila in csv.DictReader(fh):
                filas += 1
                persona = fila['Número Persona'].strip()
                if persona and persona not in padron:
                    padron[persona] = {destino: fila.get(origen, '').strip()
                                       for origen, destino in COLUMNAS_PADRON.items()}
                # Se quitan las filas IDÉNTICAS, no los pares persona+curso.
                #
                # Deduplicar por (persona, curso) parece equivalente —el motor ya
                # cuenta una vez cada par— pero no lo es: al quedarse con la
                # primera fila se tira el estatus de las demás, y con él las
                # finalizaciones completadas que venían en una repetición
                # posterior. Sobre los datos de agosto eso perdía **9,161
                # finalizaciones completadas** antes de que el motor las viera.
                #
                # La llave incluye el estatus, así que solo se van las filas que
                # de verdad sobran. Cuesta 42 mil filas y 3.5 MB más, y el motor
                # sigue contando cada par una sola vez porque hace el OR él mismo
                # (ver indiceFinalizaciones_).
                llave = (persona, clave(fila['Nombre Curso']),
                         fila['¿Lo Completó?'].strip(), fila['Sub Estatus Aprendizaje'].strip())
                if llave in vistos:
                    continue
                vistos.add(llave)
                finalizaciones.append({
                    'persona': persona,
                    'colaborador': fila['Número Colaborador'].strip(),
                    'curso': fila['Nombre Curso'].strip(),
                    'completo': fila['¿Lo Completó?'].strip(),
                    'subEstatus': fila['Sub Estatus Aprendizaje'].strip(),
                })
        print(f'  {os.path.basename(ruta):32} {filas:>9,} filas')
    return list(padron.values()), finalizaciones


# --------------------------------------------------------------------------
#  Los PDT
# --------------------------------------------------------------------------
def corrimiento(hoja, ancla='Curso'):
    """Los encabezados corridos del PDT gerencial. Ver corrimientoDeEncabezado_()."""
    encabezados = [clave(c.value) for c in hoja[1]]
    if clave(ancla) not in encabezados:
        return 0
    i = encabezados.index(clave(ancla))

    def es_rango(v):
        if isinstance(v, (datetime, date)):
            return True
        t = str('' if v is None else v).strip()
        return bool(t) and (t.replace('-', '').replace(' ', '').isdigit() or
                            t.replace('.', '', 1).isdigit())

    def es_curso(v):
        if isinstance(v, (datetime, date)):
            return False
        t = str('' if v is None else v).strip()
        return len(t) >= 4 and not es_rango(t) and any(c.isalpha() for c in t)

    en_su_lugar = a_la_derecha = filas = 0
    for fila in hoja.iter_rows(min_row=2, max_row=min(13, hoja.max_row), values_only=True):
        aqui = fila[i] if i < len(fila) else None
        al_lado = fila[i + 1] if i + 1 < len(fila) else None
        if aqui is None and al_lado is None:
            continue
        filas += 1
        if es_curso(aqui):
            en_su_lugar += 1
        if es_rango(aqui) and es_curso(al_lado):
            a_la_derecha += 1
    if filas and a_la_derecha > filas / 2 and en_su_lugar == 0:
        print(f'    aviso: "{hoja.title}" trae el encabezado corrido una columna; se corrige')
        return 1
    return 0


def leer_pestana(hoja, columnas, desplazamiento=0):
    encabezados = [clave(str(c.value or '').replace('_', ' ')) for c in hoja[1]]
    indices = {}
    for origen, destino in columnas.items():
        for alternativa in origen.split('|'):
            k = clave(alternativa.replace('_', ' '))
            if k in encabezados:
                indices[destino] = encabezados.index(k)
                break
    salida = []
    for fila in hoja.iter_rows(min_row=2, values_only=True):
        if not any(v not in (None, '') for v in fila):
            continue
        objeto = {}
        for destino, i in indices.items():
            j = i + desplazamiento
            objeto[destino] = texto_celda(fila[j] if j < len(fila) else '')
        salida.append(objeto)
    return salida


def leer_plan(ruta, familia):
    import openpyxl
    libro = openpyxl.load_workbook(ruta, data_only=True)
    print(f'  {os.path.basename(ruta):32} plan {familia}')

    def hoja_de(candidatos):
        for candidato in candidatos:
            for nombre in libro.sheetnames:
                if clave(nombre) == clave(candidato):
                    return libro[nombre]
        sys.exit(f'El plan {familia} no tiene ninguna pestaña entre {candidatos}')

    def cursos(candidatos, con_niveles):
        hoja = hoja_de(candidatos)
        corrido = corrimiento(hoja)
        base = [f for f in leer_pestana(hoja, {
            'Curso': 'curso', 'Tipo': 'tipo', 'Rango de meses para cursar': 'rango',
        }, corrido) if str(f.get('curso') or '').strip()]
        if not con_niveles:
            return base
        columnas = {'Curso': 'curso'}
        columnas.update({n: n for n in NIVELES_MATRIZ})
        matriz = leer_pestana(hoja, columnas, corrido)
        for i, fila in enumerate(base):
            marcas = matriz[i] if i < len(matriz) else {}
            niveles = {n: marcas.get(n, '') for n in NIVELES_MATRIZ}
            fila['niveles'] = niveles if any(v not in (None, '') for v in niveles.values()) else None
        return base

    especificos = leer_pestana(hoja_de(['Colaboradores_especificos', 'Colaboradores específicos']), {
        'ID': 'id', 'Puesto': 'puesto', 'Centros de costos': 'centrosDeCosto',
        'Centros que no aplican': 'noAplica', 'Centros que si aplican': 'siAplica',
    })
    lista = lambda campo: [str(f.get(campo) or '').strip()
                           for f in especificos if str(f.get(campo) or '').strip()]
    no_aplican, si_aplican = lista('noAplica'), lista('siAplica')

    return {
        'familia': familia,
        'cursosGenerales': cursos(['Cursos_asignados', 'Cursos asignados'], True),
        'cursosEspecificos': cursos(['Cursos_especificos', 'Cursos específicos'], False),
        'puestosGenerales': [f for f in leer_pestana(
            hoja_de(['Colaboradores_asignados', 'Colaboradores asignados']),
            {'ID': 'id', 'Puesto': 'puesto'}) if str(f.get('puesto') or '').strip()],
        'puestosEspecificos': [{
            'id': f.get('id'), 'puesto': f.get('puesto'),
            'centrosDeCosto': f.get('centrosDeCosto'),
            'centrosQueNoAplican': no_aplican, 'centrosQueSiAplican': si_aplican,
        } for f in especificos if str(f.get('puesto') or '').strip()],
    }


def main():
    if len(sys.argv) < 2:
        sys.exit(__doc__)
    carpeta = sys.argv[1]
    salida = sys.argv[2] if len(sys.argv) > 2 else 'fuentes.json'

    print('Finalizaciones y padrón:')
    padron, finalizaciones = leer_csv(carpeta)

    print('Planes:')
    planes = [
        leer_plan(uno(carpeta, 'PDT-operaci*n*.xlsx'), 'Colaborador'),
        leer_plan(uno(carpeta, 'PDT-gerencial*.xlsx'), 'Gerencial'),
    ]

    detalle = []
    ruta_detalle = uno(carpeta, '*detalle*colaborador*.xlsx', obligatorio=False)
    if ruta_detalle:
        import openpyxl
        hoja = openpyxl.load_workbook(ruta_detalle, data_only=True, read_only=True).worksheets[0]
        detalle = leer_pestana(hoja, {
            'NÚMERO PERSONA|Número de persona': 'numeroPersona',
            'FECHA DE INGRESO|Fecha de contratación de la empresa': 'fechaContratacion',
            'FECHA DE INGRESO DE PUESTO': 'fechaPuesto',
            'PUESTO|Nombre de puesto': 'puesto',
        })
        print(f'  {os.path.basename(ruta_detalle):32} {len(detalle):>9,} filas')

    with open(salida, 'w', encoding='utf-8') as fh:
        json.dump({'padron': padron, 'finalizaciones': finalizaciones,
                   'detalle': detalle, 'planes': planes}, fh, ensure_ascii=False)

    print(f'\n{salida}: {len(padron):,} personas · {len(finalizaciones):,} finalizaciones · '
          f'{len(detalle):,} del detalle · {os.path.getsize(salida) / 1e6:.1f} MB')


if __name__ == '__main__':
    main()
